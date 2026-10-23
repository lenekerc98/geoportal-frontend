# -*- coding: utf-8 -*-
"""
Script para incorporar la documentación completa de la App Móvil (CatastroMovil2026.apk)
dentro del Informe / Manual Técnico existente del Geoportal.
Genera el 'Manual_Tecnico_Integral_Geoportal_y_Catastro_Movil_2026.docx'.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=130, right=130):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        if b_val:
            val = b_val.get('val', 'single')
            sz = b_val.get('sz', '4')
            color = b_val.get('color', 'CCCCCC')
            b_elm = parse_xml(f'<w:{b_name} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
            tcBorders.append(b_elm)
        else:
            b_elm = parse_xml(f'<w:{b_name} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_elm)
    tcPr.append(tcBorders)

def add_callout(doc, text, title="NOTA TÉCNICA", border_color="0284C7", bg_color="F0F9FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_borders(cell, left={'val': 'single', 'sz': '24', 'color': border_color})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor.from_string(border_color)
    
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(9)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table(table, col_widths, headers, rows_data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        set_cell_borders(hdr_cells[i], bottom={'val': 'single', 'sz': '12', 'color': '0F172A'})
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.bold = True
            r.font.name = "Segoe UI"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    for r_idx, row_values in enumerate(rows_data):
        row = table.rows[r_idx + 1]
        bg_col = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.text = str(val)
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            set_cell_borders(cell, bottom={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                                   top={'val': 'single', 'sz': '4', 'color': 'E2E8F0'})
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Segoe UI"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
                
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

def merge_mobile_into_geoportal():
    src_docx = r"C:\LNCZ\proyecto-catastro-2026\informes\Informe_Esquema_Base_de_Datos_Geoportal.docx"
    dest_docx = r"C:\LNCZ\proyecto-catastro-2026\informes\Manual_Tecnico_Integral_Geoportal_y_Catastro_Movil_2026.docx"
    
    print(f"Cargando manual base del Geoportal: {src_docx}")
    doc = docx.Document(src_docx)
    
    # 1. ACTUALIZAR TÍTULO PRINCIPAL Y SUBTÍTULO
    p0 = doc.paragraphs[0]
    p0.text = "INFORME TÉCNICO INTEGRAL DE ARQUITECTURA DE SOFTWARE Y BASE DE DATOS"
    for r in p0.runs:
        r.font.name = "Segoe UI"
        r.bold = True
        r.font.size = Pt(18)
        r.font.color.rgb = RGBColor(30, 58, 138)
        
    p1 = doc.paragraphs[1]
    p1.text = "GEOPORTAL CATASTRAL 2026 & APLICACIÓN MÓVIL ANDROID DE CAMPO (CatastroMovil2026.apk)"
    for r in p1.runs:
        r.font.name = "Segoe UI"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(100, 116, 139)
        
    # 2. ENRIQUECER TABLA 0 DE METADATOS INICIALES
    t0 = doc.tables[0]
    # Agregar filas para la aplicación móvil
    new_rows = [
        ("Aplicativo Móvil Nativo:", "Catastro Móvil ECC 2026 (CatastroMovil2026.apk) | Capacitor 8.5 & React 19"),
        ("Persistencia de Campo Offline:", "IndexedDB (idb v2) + Respaldo Dual LocalStorage (Garantía Cero Pérdida 100% Offline)")
    ]
    for label, val in new_rows:
        row = t0.add_row()
        c0, c1 = row.cells[0], row.cells[1]
        c0.text = label
        c1.text = val
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=80, bottom=80, left=110, right=110)
        set_cell_margins(c1, top=80, bottom=80, left=110, right=110)
        set_cell_borders(c0, bottom={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                             top={'val': 'single', 'sz': '4', 'color': 'E2E8F0'})
        set_cell_borders(c1, bottom={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                             top={'val': 'single', 'sz': '4', 'color': 'E2E8F0'})
        p0, p1 = c0.paragraphs[0], c1.paragraphs[0]
        p0.runs[0].bold = True
        p0.runs[0].font.name = "Segoe UI"
        p0.runs[0].font.size = Pt(8.5)
        p0.runs[0].font.color.rgb = RGBColor(30, 41, 59)
        p1.runs[0].font.name = "Segoe UI"
        p1.runs[0].font.size = Pt(8.5)
        p1.runs[0].font.color.rgb = RGBColor(30, 41, 59)
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        
    # 3. AGREGAR SECCIÓN 8: ARQUITECTURA E INTEGRACIÓN DEL APLICATIVO MÓVIL
    doc.add_page_break()
    
    h1 = doc.add_heading("8. Arquitectura e Integración del Aplicativo Móvil de Campo (CatastroMovil2026.apk)", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    h1.runs[0].font.name = "Segoe UI"
    h1.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    
    p = doc.add_paragraph(
        "Catastro Móvil 2026 constituye el módulo de ingeniería de campo del Geoportal. "
        "Construido como aplicación nativa Android mediante Capacitor 8.5.2 y React 19.2, "
        "el aplicativo opera de manera completamente autónoma bajo el paradigma Offline-First. "
        "Permite a brigadistas y agrimensores capturar, delimitar, fiscalizar y registrar predios en zonas rurales "
        "o selváticas desprovistas de señal celular, garantizando la persistencia local de la información y su "
        "posterior conciliación y sincronización con la base central de datos PostgreSQL/PostGIS."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    for r in p.runs:
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)
        
    add_callout(
        doc,
        "El ejecutable oficial CatastroMovil2026.apk reside en C:\\LNCZ\\proyecto-catastro-2026\\movil\\CatastroMovil2026.apk. "
        "El entorno de compilación es 100% portable y no requiere instalación de Android Studio gracias al script build_apk.js "
        "y sus herramientas empaquetadas en movil/tools/.",
        title="DISTRIBUCIÓN OFICIAL DEL APLICATIVO MÓVIL",
        border_color="0284C7",
        bg_color="F0F9FF"
    )

    # 8.1 Ficha Técnica y Runtime
    h2 = doc.add_heading("8.1. Ficha Técnica y Runtime Nativo Android", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.runs[0].font.name = "Segoe UI"
    h2.runs[0].font.color.rgb = RGBColor(14, 116, 144)

    t_apk_spec = doc.add_table(rows=8, cols=3)
    format_table(
        t_apk_spec,
        col_widths=[1.6, 1.8, 3.4],
        headers=["Componente / Capa", "Especificación Técnica", "Propósito en el Ecosistema Catastral"],
        rows_data=[
            ("Shell Nativo", "Capacitor 8.5.2 / Android Gradle", "Puente bidireccional hacia APIs de hardware: GPS de alta precisión, cámara y almacenamiento."),
            ("Identificador de Paquete", "com.gad.catastromovil.ecc", "Namespace único institucional registrado para distribución municipal."),
            ("Versión de Producción", "v4.7.6 | versionCode: 46", "Compilación de producción sincronizada con la base de datos oficial 2026."),
            ("Compatibilidad Android", "minSdkVersion: 24 | targetSdkVersion: 34", "Opera en el 98% de dispositivos Android activos en el mercado (Android 7.0 hasta Android 14+)."),
            ("Motor Cartográfico", "Leaflet 1.9.4 & React-Leaflet 5.0", "Visualizador ligero capaz de desplegar ortofotos Google Hybrid, Esri World Imagery y OpenStreetMap."),
            ("Motor Geodésico", "Proj4.js 2.20", "Conversión en tiempo real de coordenadas geográficas GPS a coordenadas planas UTM 17S Ecuador."),
            ("Motor de Interfaz", "React 19.2 + Vite 8.1", "Renderizado acelerado por hardware a 60 FPS con diseño ergonómico para una mano en campo.")
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 8.2 Persistencia Local y Garantía Cero Pérdida
    h2 = doc.add_heading("8.2. Persistencia Local Offline-First y Garantía 'Cero Pérdida'", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.runs[0].font.name = "Segoe UI"
    h2.runs[0].font.color.rgb = RGBColor(14, 116, 144)

    p = doc.add_paragraph(
        "En misiones de campo de alta complejidad topográfica, la pérdida accidental de datos es inaceptable. "
        "El servicio mobileDB.js implementa una base IndexedDB 'catastro_movil_db' (Versión 2) con una arquitectura "
        "de doble escritura simultánea (Dual-Write) en IndexedDB y LocalStorage. Los Object Stores se articulan de la siguiente manera:"
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)

    t_stores = doc.add_table(rows=5, cols=3)
    format_table(
        t_stores,
        col_widths=[1.8, 1.5, 3.5],
        headers=["Object Store", "Clave Primaria (keyPath)", "Función de Almacenamiento"],
        rows_data=[
            ("offline_predios", "offline_id (Índices: sync_status, created_at)", "Predios levantados en campo con estado 'pending', 'syncing', 'synced' o 'error'. Incluye vértices UTM, fotos y linderos."),
            ("cached_predios", "id (ID del servidor)", "Predios existentes en la base central descargados para superposición y consulta visual en el mapa."),
            ("cached_posesionarios", "key (Índice: cedula)", "Directorio de posesionarios de la entidad para autocompletado en la ficha móvil en 0 milisegundos."),
            ("settings", "key", "Configuración del entorno: URL del servidor, token JWT, UUID del dispositivo móvil y modo de sincronización.")
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 8.3 Motor Matemático y Topográfico Planar
    h2 = doc.add_heading("8.3. Motor Matemático y Topográfico Planar en Campo (EPSG:32717)", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.runs[0].font.name = "Segoe UI"
    h2.runs[0].font.color.rgb = RGBColor(14, 116, 144)

    p = doc.add_paragraph(
        "Catastro Móvil integra en geoUtils.js los mismos modelos de cómputo del Geoportal central, garantizando exactitud matemática:\n"
        "1. Proyección Cartográfica: Transforma lat/lng satelitales al sistema oficial EPSG:32717 (UTM Zona 17 Sur Ecuador) mediante Proj4.\n"
        "2. Cálculo de Superficie (Shoelace Planar): Aplica la fórmula del área poligonal por producto cruzado sobre coordenadas planas métricas (m² y ha).\n"
        "3. Normalización Oficial (Vértice NW y Sentido Horario): La función normalizeVerticesClockwiseFromNW() garantiza que todo polígono inicie "
        "en el vértice más al Nor-Oeste (P01) y recorra en estricto sentido horario (P02, P03... Pn), acatando los estándares de la Dirección de Catastros.\n"
        "4. Cálculo de Acimut y Rumbo: La función calculateRumbo() traduce el azimut geodésico a rumbos cuadrantales (ej: N 42° 15' E)."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    for r in p.runs:
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)

    # 8.4 Enlace con el Geoportal REST API y Conciliación
    h2 = doc.add_heading("8.4. Enlace con el Geoportal REST API y Protocolo de Sincronización", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.runs[0].font.name = "Segoe UI"
    h2.runs[0].font.color.rgb = RGBColor(14, 116, 144)

    t_api = doc.add_table(rows=6, cols=3)
    format_table(
        t_api,
        col_widths=[1.8, 1.4, 3.6],
        headers=["Endpoint del Geoportal", "Método HTTP", "Operación y Tratamiento Móvil"],
        rows_data=[
            ("/api/token", "POST (Form)", "Autenticación de técnicos con credenciales institucionales y emisión de JWT."),
            ("/api/login-brigadista", "POST (JSON)", "Registro ágil de operadores temporales con Nombre, Apellido y UUID de hardware."),
            ("/api/gis/predios", "GET / POST", "Descarga de cartografía en lote y subida de nuevos predios vectoriales en GeoJSON."),
            ("/api/gis/predios/verificar-conflicto", "POST (JSON)", "Detección proactiva de colisiones territoriales por clave catastral preexistente."),
            ("/api/gis/posesionarios/buscar/{cedula}", "GET", "Búsqueda en servidor de ciudadanos para vinculación al predio levantado.")
        ]
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_callout(
        doc,
        "Si al sincronizar un predio la clave catastral ya existe en el servidor central, el APK presenta una alerta interactiva "
        "SweetAlert2 con tres alternativas: 1) Sobreescribir el predio existente, 2) Conservar el registro del servidor y marcar como omitido, "
        "o 3) Asignar en el acto un nuevo código catastral y reenviar.",
        title="CONCILIACIÓN INTERACTIVA DE CONFLICTOS",
        border_color="F59E0B",
        bg_color="FEFCE8"
    )

    # 8.5 Manual de Operación en Campo
    h2 = doc.add_heading("8.5. Manual de Operación en Campo (Flujo de Trabajo para Brigadistas)", level=2)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.runs[0].font.name = "Segoe UI"
    h2.runs[0].font.color.rgb = RGBColor(14, 116, 144)

    steps_op = [
        "1. Preparación en Oficina: Abrir el APK con Wi-Fi, iniciar sesión y en Ajustes presionar 'Descargar Predios del Servidor' para precargar la base y fotos satelitales.",
        "2. Llegada a Campo: Abrir la pestaña 'Mapa' y constatar que el indicador superior muestre 'GPS Fix' en color verde (< 5 m de precisión).",
        "3. Modo Caminar (Levantamiento Perimétrico): Abrir la barra de dibujo, presionar 'Caminar Perímetro', situarse en la primera esquina y pulsar '+ Capturar Vértice'. Caminar por el lindero hasta la siguiente esquina y capturar P02, repitiendo en sentido horario hasta cerrar el lote.",
        "4. Generación de Ficha: Pulsar 'Finalizar y Crear Ficha'. El sistema computa área y perímetro exactos, desglosa los tramos métricos y solicita ingresar código catastral, cédula del posesionario (con búsqueda en 0 ms) y colindantes.",
        "5. Emisión de Certificado Planimétrico Móvil: Desde la ficha o tocando el predio en mapa, se puede visualizar la lámina técnica con cuadrícula UTM oficial, cuadro de rumbos y distancias y código QR institucional.",
        "6. Sincronización al Servidor: Al regresar a zona con conectividad, dirigirse a 'Sincronizar' y presionar 'Sincronizar Todos'."
    ]
    for step in steps_op:
        p_step = doc.add_paragraph(step)
        p_step.paragraph_format.line_spacing = 1.15
        p_step.paragraph_format.space_after = Pt(4)
        for r in p_step.runs:
            r.font.name = "Segoe UI"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(30, 41, 59)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Guardar documento resultante
    doc.save(dest_docx)
    print(f"Manual integral guardado exitosamente en:\n{dest_docx}")

if __name__ == '__main__':
    merge_mobile_into_geoportal()
